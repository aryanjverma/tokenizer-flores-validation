"""One slide: same meaning, fake token splits — English clean, Hindi/Hungarian shredded."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "fertility_text_example_slide.pptx"

NAVY = RGBColor(0x0F, 0x2A, 0x44)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x67, 0x72)
CREAM = RGBColor(0xF7, 0xF4, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ENG_C = RGBColor(0x2A, 0x9D, 0x8F)
HIN_C = RGBColor(0xE0, 0x6C, 0x4F)
HUN_C = RGBColor(0xE9, 0xA8, 0x2C)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def textbox(slide, left, top, width, height, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return box


def chip(slide, left, top, label, color, text_color=WHITE, height=Inches(0.42)):
    # width from label length
    w = Inches(max(0.55, 0.11 * len(label) + 0.35))
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, height)
    fill(sh, color)
    try:
        sh.adjustments[0] = 0.25
    except Exception:
        pass
    tf = sh.text_frame
    tf.word_wrap = False
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = label
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = text_color
    r.font.name = "Calibri"
    # vertical-ish via margin
    tf.margin_top = Pt(6)
    return w


def row_of_chips(slide, left, top, tokens, color, max_right=Inches(12.8)):
    x = left
    y = top
    row_h = Inches(0.5)
    gap = Inches(0.08)
    for tok in tokens:
        w = Inches(max(0.55, 0.11 * len(tok) + 0.35))
        if x + w > max_right:
            x = left
            y += row_h
        chip(slide, x, y, tok, color)
        x += w + gap
    return y + row_h


def language_block(slide, top, lang, gloss, surface, tokens, color, n_label):
    # header
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), top, Inches(1.7), Inches(0.38))
    fill(badge, color)
    try:
        badge.adjustments[0] = 0.3
    except Exception:
        pass
    tf = badge.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = lang
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = "Calibri"
    tf.margin_top = Pt(4)

    textbox(slide, Inches(2.3), top, Inches(8.5), Inches(0.38), surface, size=15, bold=True, color=INK)
    textbox(slide, Inches(11.0), top, Inches(1.9), Inches(0.38), n_label, size=14, bold=True, color=color, align=PP_ALIGN.RIGHT)

    y = row_of_chips(slide, Inches(0.55), top + Inches(0.5), tokens, color)
    return y + Inches(0.15)


def main():
    # Same meaning. Fake / pedagogical splits — not real BPE.
    # English: clean word-ish pieces
    # Hindi / Hungarian: shredded into tiny fragments (illustrative "bad tokenization")
    eng_surface = "The children are playing in the park."
    hin_surface = "बच्चे पार्क में खेल रहे हैं।"
    hun_surface = "A gyerekek a parkban játszanak."

    eng_toks = ["The", "children", "are", "playing", "in", "the", "park", "."]
    hin_toks = ["ब", "च्", "चे", "पा", "र्", "क", "मे", "ं", "खे", "ल", "र", "हे", "है", "ं", "।"]
    hun_toks = ["A", "gyer", "ek", "ek", "a", "park", "ban", "ját", "sz", "anak", "."]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.95))
    fill(bar, NAVY)
    textbox(s, Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.5),
            "Same sentence. Different token bills.", size=26, bold=True, color=WHITE)
    textbox(s, Inches(0.45), Inches(1.05), Inches(12.4), Inches(0.35),
            "Illustrative splits (not a real tokenizer) — fertility = how many pieces the text breaks into.",
            size=13, color=MUTED)

    y = Inches(1.55)
    y = language_block(
        s, y, "English", "eng", eng_surface, eng_toks, ENG_C,
        f"{len(eng_toks)} tokens",
    )
    y = language_block(
        s, y + Inches(0.1), "Hindi", "hin", hin_surface, hin_toks, HIN_C,
        f"{len(hin_toks)} tokens",
    )
    y = language_block(
        s, y + Inches(0.1), "Hungarian", "hun", hun_surface, hun_toks, HUN_C,
        f"{len(hun_toks)} tokens",
    )

    note = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.7))
    fill(note, CREAM)
    textbox(s, Inches(0.7), Inches(6.7), Inches(12), Inches(0.45),
            "Takeaway: morphologically rich / non-Latin text often becomes many tiny tokens → higher cost, less context, worse accuracy.",
            size=14, color=INK)

    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
